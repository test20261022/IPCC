"""名单文件解析与行级校验（OB-011 / §24）。

纯函数实现（不依赖数据库），便于用单测覆盖各种表头、多号码与异常数据。

支持：
- 格式：.xlsx / .csv（CSV 自动识别 UTF-8 / UTF-8-BOM / GBK）
- 动态表头：核心列按别名识别，其余列全部进 dynamic_fields（JSONB）
- 多号码列：所有号码类列按列顺序合并为一个号码清单
"""

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.exceptions import ParamError
from app.utils.validators import is_valid_phone

SUPPORTED_EXTENSIONS = (".xlsx", ".csv")

#: 核心列别名（比较前会先去空格/下划线/横线并转小写）
_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "user_id": ("userid", "用户id", "用户编号", "客户id", "客户编号", "会员id", "会员编号"),
    "order_id": ("orderid", "订单号", "订单id", "订单编号", "下单单号", "主订单号"),
    "recipient_name": ("收件人", "收件人姓名", "姓名", "客户姓名", "联系人", "收货人"),
}

#: 号码类列别名（可多列，全部作为被叫号码）
_PHONE_ALIASES = ("电话", "号码", "手机", "phone", "mobile", "tel")

#: 「收件人号码」专用别名，命中则该列值写入 recipient_phone
_RECIPIENT_PHONE_ALIASES = ("收件人号码", "收件号码", "收件人电话", "收货人电话")

_HEADER_NOISE = re.compile(r"[\s_\-（）()·、]")
_EMPTY_ROW_MARKERS = ("", None)


def normalize_header(value: Any) -> str:
    """表头归一化：去空白与常见分隔符后转小写。"""
    text = "" if value is None else str(value)
    return _HEADER_NOISE.sub("", text).strip().lower()


@dataclass
class HeaderMapping:
    """表头映射结果。"""

    fields: dict[str, int] = field(default_factory=dict)
    phone_indexes: list[int] = field(default_factory=list)
    recipient_phone_index: int | None = None
    dynamic_indexes: dict[int, str] = field(default_factory=dict)
    header_names: list[str] = field(default_factory=list)

    @property
    def recognized(self) -> bool:
        return bool(self.fields or self.phone_indexes)


@dataclass
class ParsedListRow:
    """通过校验的名单行。"""

    row_no: int
    user_id: str | None
    order_id: str | None
    recipient_name: str | None
    recipient_phone: str | None
    phone_numbers: list[str]
    dynamic_fields: dict[str, Any]


@dataclass
class RowFailure:
    """行级失败明细（§25：行号 / 字段 / 原因 / 原始值）。"""

    row: int
    field: str
    reason: str
    raw: str | None


def detect_format(filename: str) -> str:
    """按扩展名判定格式；不支持的格式直接拒绝。"""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ParamError("仅支持 .xlsx 与 .csv 格式的名单文件")
    return "xlsx" if suffix == ".xlsx" else "csv"


def read_rows(content: bytes, file_format: str) -> list[list[str]]:
    """读取文件为二维字符串数组（含表头行）。"""
    return _read_xlsx(content) if file_format == "xlsx" else _read_csv(content)


def _read_xlsx(content: bytes) -> list[list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - 统一转成业务错误
        raise ParamError("Excel 文件解析失败，请确认文件未损坏且为 .xlsx 格式") from exc

    try:
        sheet = workbook.active
        return [
            ["" if cell is None else str(cell).strip() for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()


def _read_csv(content: bytes) -> list[list[str]]:
    text = _decode_csv(content)
    return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text))]


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParamError("CSV 编码无法识别，请另存为 UTF-8 或 GBK 后重试")


def parse_header(rows: list[list[str]]) -> HeaderMapping:
    """识别表头：核心列按别名匹配，其余非空表头作为动态字段。"""
    if not rows:
        raise ParamError("名单文件内容为空")
    if len(rows) < 2:
        raise ParamError("名单文件只有表头，没有可导入的数据行")

    header = rows[0]
    mapping = HeaderMapping(header_names=list(header))

    for index, raw in enumerate(header):
        normalized = normalize_header(raw)
        if not normalized:
            continue

        matched_field = next(
            (
                name
                for name, aliases in _HEADER_ALIASES.items()
                if normalized in aliases and name not in mapping.fields
            ),
            None,
        )
        if matched_field:
            mapping.fields[matched_field] = index
            continue

        if any(alias in normalized for alias in _RECIPIENT_PHONE_ALIASES):
            if mapping.recipient_phone_index is None:
                mapping.recipient_phone_index = index
            mapping.phone_indexes.append(index)
            continue

        if any(alias in normalized for alias in _PHONE_ALIASES):
            mapping.phone_indexes.append(index)
            continue

        mapping.dynamic_indexes[index] = str(raw).strip()

    if not mapping.recognized:
        raise ParamError(
            "未能识别名单表头，请至少包含 UserID / OrderID / 收件人姓名 / 电话号 之一"
        )
    return mapping


def parse_rows(
    rows: list[list[str]], mapping: HeaderMapping
) -> tuple[list[ParsedListRow], list[RowFailure]]:
    """逐行校验，返回 (可用行, 失败明细)。行号按 Excel 视角从 1 开始（表头为第 1 行）。"""
    parsed: list[ParsedListRow] = []
    failures: list[RowFailure] = []

    for offset, row in enumerate(rows[1:], start=2):
        if not _has_value(row):
            continue

        phones: list[str] = []
        invalid = False
        for index in mapping.phone_indexes:
            if index >= len(row):
                continue
            value = row[index].strip()
            if not value:
                continue
            if not is_valid_phone(value):
                failures.append(
                    RowFailure(
                        row=offset,
                        field=_header_name(mapping, index),
                        reason="号码格式不正确",
                        raw=value,
                    )
                )
                invalid = True
                continue
            phones.append(value)

        if invalid:
            continue

        user_id = _cell(row, mapping.fields.get("user_id"))
        order_id = _cell(row, mapping.fields.get("order_id"))
        if not phones and not user_id and not order_id:
            failures.append(
                RowFailure(
                    row=offset,
                    field="",
                    reason="缺少 UserID / OrderID / 号码，无法外呼",
                    raw=" | ".join(row)[:255],
                )
            )
            continue

        recipient_phone = _cell(row, mapping.recipient_phone_index)
        if recipient_phone is None and phones:
            recipient_phone = phones[0]

        dynamic_fields: dict[str, Any] = {}
        for index, name in mapping.dynamic_indexes.items():
            value = _cell(row, index)
            if value:
                dynamic_fields[name] = value

        parsed.append(
            ParsedListRow(
                row_no=offset,
                user_id=user_id,
                order_id=order_id,
                recipient_name=_cell(row, mapping.fields.get("recipient_name")),
                recipient_phone=recipient_phone,
                phone_numbers=phones,
                dynamic_fields=dynamic_fields,
            )
        )

    return parsed, failures


def _cell(row: list[str], index: int | None) -> str | None:
    """取指定列的值，空值返回 None。"""
    if index is None or index >= len(row):
        return None
    value = row[index].strip()
    return value or None


def _header_name(mapping: HeaderMapping, index: int) -> str:
    if index < len(mapping.header_names):
        return str(mapping.header_names[index]).strip()
    return ""


def _has_value(row: list[str]) -> bool:
    return any(cell not in _EMPTY_ROW_MARKERS and str(cell).strip() for cell in row)
