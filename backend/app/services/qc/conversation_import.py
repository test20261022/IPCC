"""对话批量导入：解析 TXT / JSON / XLSX 并落地为对话与消息（§24 同类口径）。

支持三种格式：
- JSON：顶层数组或 {"items": [...]}，每项含对话元信息与 messages 列表；
- TXT：一行一条对话，制表符分隔「对话时间 | 客服 | 客户标识 | 对话内容」，
  对话内容按 客服: / 客户: 前缀切分说话人（兼容中文冒号）；
- XLSX：表头 started_at / agent_name / agent_ext / customer_id / role / content，
  同一 ext_ref_no（或同 started_at+agent_name）的多行合并为一条对话的多条消息。

幂等（§78）：同租户下 ext_ref_no 已存在则跳过并计入失败明细，不中断整批。
"""

import io
import json
import logging
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.exceptions import ParamError
from app.models.qc import QcImportBatch
from app.models.qc.enums import (
    ConversationChannel,
    ConversationSource,
    MessageRole,
    MessageSource,
    QcImportFormat,
    QcImportStatus,
    TranscriptStatus,
)
from app.repositories.qc.conversation_repository import QcConversationRepository
from app.utils.storage import save_bytes

logger = logging.getLogger(__name__)

#: 扩展名 → 导入格式
SUPPORTED_EXTENSIONS: dict[str, QcImportFormat] = {
    ".json": QcImportFormat.JSON,
    ".txt": QcImportFormat.TXT,
    ".xlsx": QcImportFormat.XLSX,
}

#: XLSX 约定表头（可选列缺失时跳过）
XLSX_HEADERS = (
    "ext_ref_no",
    "started_at",
    "agent_name",
    "agent_ext",
    "customer_id",
    "role",
    "content",
)

_ROLE_ALIASES: dict[str, MessageRole] = {
    "agent": MessageRole.AGENT,
    "客服": MessageRole.AGENT,
    "坐席": MessageRole.AGENT,
    "customer": MessageRole.CUSTOMER,
    "客户": MessageRole.CUSTOMER,
    "用户": MessageRole.CUSTOMER,
    "system": MessageRole.SYSTEM,
    "系统": MessageRole.SYSTEM,
}

_CHANNEL_ALIASES: dict[str, ConversationChannel] = {
    "chat": ConversationChannel.CHAT,
    "文本": ConversationChannel.CHAT,
    "在线": ConversationChannel.CHAT,
    "voice": ConversationChannel.VOICE,
    "语音": ConversationChannel.VOICE,
    "电话": ConversationChannel.VOICE,
    "other": ConversationChannel.OTHER,
    "其他": ConversationChannel.OTHER,
}

#: TXT 说话人前缀（兼容中文冒号），切分出发言人段落
_SPEAKER_SPLIT = re.compile(r"(客服|坐席|客户|用户|系统)\s*[:：]")
_SPEAKER_ROLES: dict[str, MessageRole] = {
    "客服": MessageRole.AGENT,
    "坐席": MessageRole.AGENT,
    "客户": MessageRole.CUSTOMER,
    "用户": MessageRole.CUSTOMER,
    "系统": MessageRole.SYSTEM,
}

_HEADER_NOISE = re.compile(r"[\s_\-]")

_DATETIME_FORMATS = (
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%Y/%m/%d %H:%M:%S",
    "%Y%m%d%H%M%S",
    "%Y%m%d",
)


@dataclass
class ParsedMessage:
    """解析出的一条消息。"""

    role: MessageRole
    content: str
    start_ms: int | None = None
    end_ms: int | None = None


@dataclass
class ParsedConversation:
    """解析出的一条对话（含消息列表）。"""

    row_no: int
    started_at: datetime
    agent_name: str = ""
    ext_ref_no: str | None = None
    agent_ext: str | None = None
    customer_id: str | None = None
    channel: ConversationChannel = ConversationChannel.CHAT
    ended_at: datetime | None = None
    duration: int = 0
    recording_path: str | None = None
    messages: list[ParsedMessage] = field(default_factory=list)


@dataclass
class ImportFailure:
    """一行 / 一条对话的失败明细。"""

    row: int
    reason: str
    raw: str | None = None


# ---------------- 解析（纯函数） ----------------

def detect_format(filename: str) -> QcImportFormat:
    """按扩展名判定格式；不支持的格式直接拒绝。"""
    file_format = SUPPORTED_EXTENSIONS.get(Path(filename).suffix.lower())
    if file_format is None:
        raise ParamError("仅支持 .json / .txt / .xlsx 格式的对话文件")
    return file_format


def parse_file(
    content: bytes, file_format: QcImportFormat
) -> tuple[list[ParsedConversation], list[ImportFailure]]:
    """解析文件为 (对话列表, 失败明细)。"""
    if file_format is QcImportFormat.JSON:
        return _parse_json(content)
    if file_format is QcImportFormat.TXT:
        return _parse_txt(content)
    return _parse_xlsx(content)


def _parse_json(content: bytes) -> tuple[list[ParsedConversation], list[ImportFailure]]:
    try:
        payload = json.loads(_decode_text(content))
    except ValueError as exc:
        raise ParamError("JSON 文件解析失败，请确认内容为合法 JSON") from exc

    items = payload.get("items") if isinstance(payload, dict) else payload
    if not isinstance(items, list):
        raise ParamError("JSON 顶层应为数组，或包含 items 数组的对象")

    parsed: list[ParsedConversation] = []
    failures: list[ImportFailure] = []
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            failures.append(ImportFailure(row=index, reason="对话项不是对象", raw=str(item)[:255]))
            continue
        started_at = _parse_datetime(item.get("started_at"))
        if started_at is None:
            failures.append(
                ImportFailure(row=index, reason="缺少或无法解析 started_at", raw=_dump(item))
            )
            continue
        parsed.append(
            ParsedConversation(
                row_no=index,
                started_at=started_at,
                agent_name=str(item.get("agent_name") or "").strip(),
                ext_ref_no=_text(item.get("ext_ref_no")),
                agent_ext=_text(item.get("agent_ext")),
                customer_id=_text(item.get("customer_id")),
                channel=_coerce_channel(item.get("channel")),
                ended_at=_parse_datetime(item.get("ended_at")),
                duration=_to_int(item.get("duration"), 0) or 0,
                recording_path=_text(item.get("recording_path")),
                messages=_parse_json_messages(item.get("messages")),
            )
        )
    return parsed, failures


def _parse_json_messages(value: Any) -> list[ParsedMessage]:
    if not isinstance(value, list):
        return []
    messages: list[ParsedMessage] = []
    for entry in value:
        if not isinstance(entry, dict):
            continue
        messages.append(
            ParsedMessage(
                role=_coerce_role(entry.get("role")),
                content=str(entry.get("content") or ""),
                start_ms=_to_int(entry.get("start_ms")),
                end_ms=_to_int(entry.get("end_ms")),
            )
        )
    return messages


def _parse_txt(content: bytes) -> tuple[list[ParsedConversation], list[ImportFailure]]:
    text = _decode_text(content)
    parsed: list[ParsedConversation] = []
    failures: list[ImportFailure] = []

    for offset, raw_line in enumerate(text.splitlines(), start=1):
        if not raw_line.strip():
            continue
        parts = [part.strip() for part in raw_line.split("\t")]
        if len(parts) < 4:
            failures.append(
                ImportFailure(
                    row=offset,
                    reason="字段不足（需 4 列：对话时间/客服/客户标识/对话内容）",
                    raw=raw_line.strip()[:255],
                )
            )
            continue
        started_at = _parse_datetime(parts[0])
        if started_at is None:
            failures.append(
                ImportFailure(row=offset, reason="无法解析对话时间", raw=raw_line.strip()[:255])
            )
            continue
        parsed.append(
            ParsedConversation(
                row_no=offset,
                started_at=started_at,
                agent_name=parts[1],
                customer_id=parts[2] or None,
                messages=_split_speakers("\t".join(parts[3:])),
            )
        )
    return parsed, failures


def _split_speakers(content: str) -> list[ParsedMessage]:
    """按 客服: / 客户: 前缀切分说话人；无前缀时整段作为一条客户消息。"""
    text = content.strip()
    if not text:
        return []
    matches = list(_SPEAKER_SPLIT.finditer(text))
    if not matches:
        return [ParsedMessage(role=MessageRole.CUSTOMER, content=text)]

    messages: list[ParsedMessage] = []
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        body = text[start:end].strip()
        if not body:
            continue
        messages.append(
            ParsedMessage(
                role=_SPEAKER_ROLES.get(match.group(1), MessageRole.CUSTOMER),
                content=body,
            )
        )
    return messages or [ParsedMessage(role=MessageRole.CUSTOMER, content=text)]


def _parse_xlsx(content: bytes) -> tuple[list[ParsedConversation], list[ImportFailure]]:
    rows = _read_xlsx(content)
    if not rows:
        raise ParamError("Excel 文件内容为空")

    header = [_normalize_header(cell) for cell in rows[0]]
    # 表头单元格已归一化（去空格/下划线/连字符），因此查找键也必须归一化后再比对
    mapping = {
        name: header.index(_normalize_header(name))
        for name in XLSX_HEADERS
        if _normalize_header(name) in header
    }
    if "started_at" not in mapping or "content" not in mapping:
        raise ParamError("Excel 表头需至少包含 started_at 与 content 列")

    groups: dict[str, ParsedConversation] = {}
    failures: list[ImportFailure] = []

    for offset, row in enumerate(rows[1:], start=2):
        if not _row_has_value(row):
            continue
        started_raw = _cell(row, mapping.get("started_at"))
        ext_ref = _cell(row, mapping.get("ext_ref_no"))
        agent_name = _cell(row, mapping.get("agent_name")) or ""
        key = ext_ref or f"{started_raw}|{agent_name}"

        group = groups.get(key)
        if group is None:
            started_at = _parse_datetime(started_raw)
            if started_at is None:
                failures.append(
                    ImportFailure(row=offset, reason="缺少或无法解析 started_at", raw=_raw(row))
                )
                continue
            group = ParsedConversation(
                row_no=offset,
                started_at=started_at,
                agent_name=agent_name,
                ext_ref_no=ext_ref,
                agent_ext=_cell(row, mapping.get("agent_ext")),
                customer_id=_cell(row, mapping.get("customer_id")),
            )
            groups[key] = group

        message_content = _cell(row, mapping.get("content"))
        if message_content:
            group.messages.append(
                ParsedMessage(
                    role=_coerce_role(_cell(row, mapping.get("role"))),
                    content=message_content,
                )
            )
    return list(groups.values()), failures


def _read_xlsx(content: bytes) -> list[list[Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - 统一转成业务错误
        raise ParamError("Excel 文件解析失败，请确认文件未损坏且为 .xlsx 格式") from exc
    try:
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


# ---------------- 解析辅助 ----------------

def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParamError("文件编码无法识别，请另存为 UTF-8 或 GBK 后重试")


def _parse_datetime(value: Any) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        pass
    for fmt in _DATETIME_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _coerce_role(value: Any) -> MessageRole:
    if isinstance(value, MessageRole):
        return value
    return _ROLE_ALIASES.get(str(value or "").strip().lower(), MessageRole.CUSTOMER)


def _coerce_channel(value: Any) -> ConversationChannel:
    if isinstance(value, ConversationChannel):
        return value
    return _CHANNEL_ALIASES.get(str(value or "").strip().lower(), ConversationChannel.CHAT)


def _to_int(value: Any, default: int | None = None) -> int | None:
    if value is None or value == "":
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _normalize_header(value: Any) -> str:
    return _HEADER_NOISE.sub("", str(value or "")).strip().lower()


def _cell(row: list[Any], index: int | None) -> str | None:
    if index is None or index >= len(row):
        return None
    return _text(row[index])


def _row_has_value(row: list[Any]) -> bool:
    return any(cell is not None and str(cell).strip() for cell in row)


def _dump(value: Any) -> str:
    try:
        return json.dumps(value, ensure_ascii=False)[:255]
    except (TypeError, ValueError):
        return str(value)[:255]


def _raw(row: list[Any]) -> str:
    return " | ".join(str(cell) for cell in row if cell is not None)[:255]


# ---------------- 导入服务 ----------------

class QcConversationImportService:
    """对话导入业务。事务在本层提交。"""

    def __init__(self, db: Session) -> None:
        self.db = db
        self.repo = QcConversationRepository(db)

    def import_file(self, *, tenant_id: int, filename: str, content: bytes) -> QcImportBatch:
        """解析并导入对话，返回导入批次。"""
        self._validate_file(filename, content)
        file_format = detect_format(filename)
        parsed, failures = parse_file(content, file_format)

        if len(parsed) + len(failures) > settings.QC_IMPORT_MAX_ROWS:
            raise ParamError(f"导入条数超过上限（{settings.QC_IMPORT_MAX_ROWS} 条）")
        if not parsed and not failures:
            raise ParamError("文件中没有有效数据")

        kept, duplicate_failures = self._drop_duplicates(tenant_id, parsed)
        all_failures = [*failures, *duplicate_failures]

        relative_path = save_bytes(
            subdir="qc_conversations", filename=filename, content=content
        )
        batch = self.repo.create_batch(
            tenant_id=tenant_id,
            file_name=filename,
            file_path=relative_path,
            format=file_format,
            total_rows=len(parsed) + len(failures),
            success_rows=0,
            failed_rows=len(all_failures),
            status=QcImportStatus.IMPORTING,
            error_detail=None,
        )
        # 先提交批次，保证即使后续写入失败也留有审计记录
        self.db.commit()
        batch_id = batch.id

        try:
            inserted = self._write_conversations(tenant_id, kept)
        except Exception:
            self._mark_batch_failed(batch_id)
            raise

        batch = self.repo.get_batch(batch_id)
        if batch is None:  # pragma: no cover - 提交后必然存在
            raise ParamError("导入批次写入失败")
        batch.success_rows = inserted
        batch.failed_rows = len(all_failures)
        batch.status = self._resolve_status(inserted, len(all_failures))
        batch.error_detail = self._dump_failures(all_failures)
        self.db.commit()

        logger.info(
            "对话导入完成：tenant=%s batch=%s 总=%s 成功=%s 失败=%s",
            tenant_id,
            batch_id,
            batch.total_rows,
            inserted,
            len(all_failures),
        )
        return batch

    def list_batches(
        self, *, page: int, page_size: int, tenant_id: int | None = None
    ) -> tuple[list[QcImportBatch], int]:
        items = self.repo.list_batches_paged(page=page, page_size=page_size, tenant_id=tenant_id)
        total = self.repo.count_batches(tenant_id=tenant_id)
        return items, total

    # ---------- 内部 ----------

    @staticmethod
    def _validate_file(filename: str, content: bytes) -> None:
        if not filename:
            raise ParamError("请选择要导入的对话文件")
        if not content:
            raise ParamError("对话文件内容为空")

    def _drop_duplicates(
        self, tenant_id: int, parsed: list[ParsedConversation]
    ) -> tuple[list[ParsedConversation], list[ImportFailure]]:
        """同租户 ext_ref_no 已存在的对话跳过并计入失败明细（幂等）。"""
        refs = [item.ext_ref_no for item in parsed if item.ext_ref_no]
        existing = self.repo.list_existing_ext_refs(tenant_id, list(refs))

        kept: list[ParsedConversation] = []
        duplicates: list[ImportFailure] = []
        for item in parsed:
            if item.ext_ref_no and item.ext_ref_no in existing:
                duplicates.append(
                    ImportFailure(row=item.row_no, reason="已存在", raw=item.ext_ref_no)
                )
                continue
            kept.append(item)
            if item.ext_ref_no:
                existing.add(item.ext_ref_no)
        return kept, duplicates

    def _write_conversations(
        self, tenant_id: int, conversations: list[ParsedConversation]
    ) -> int:
        """批量写入对话与其消息，返回写入的对话数。"""
        if not conversations:
            return 0

        values = [self._to_conversation_values(tenant_id, item) for item in conversations]
        conversation_ids = self.repo.bulk_insert_conversations(values)

        message_values: list[dict[str, object]] = []
        for conversation_id, item in zip(conversation_ids, conversations, strict=True):
            for seq, message in enumerate(item.messages, start=1):
                message_values.append(
                    {
                        "tenant_id": tenant_id,
                        "conversation_id": conversation_id,
                        "seq": seq,
                        "role": message.role,
                        "content": message.content,
                        "start_ms": message.start_ms,
                        "end_ms": message.end_ms,
                        "source": MessageSource.TEXT,
                    }
                )
        self.repo.bulk_insert_messages(message_values)
        return len(conversation_ids)

    @staticmethod
    def _to_conversation_values(tenant_id: int, item: ParsedConversation) -> dict[str, object]:
        return {
            "tenant_id": tenant_id,
            "ext_ref_no": item.ext_ref_no,
            "agent_name": item.agent_name,
            "agent_ext": item.agent_ext,
            "customer_id": item.customer_id,
            "channel": item.channel,
            "source": ConversationSource.IMPORT,
            "started_at": item.started_at,
            "ended_at": item.ended_at,
            "duration": item.duration,
            # msg_count 直接按实际写入的消息数回填，避免写入后再次逐行更新
            "msg_count": len(item.messages),
            "recording_path": item.recording_path,
            "transcript_status": _import_transcript_status(item),
            "transcript_engine": None,
            "is_deleted": False,
        }

    @staticmethod
    def _resolve_status(success: int, failed: int) -> QcImportStatus:
        if failed == 0:
            return QcImportStatus.SUCCESS
        if success == 0:
            return QcImportStatus.FAILED
        return QcImportStatus.PARTIAL_FAILED

    @staticmethod
    def _dump_failures(failures: list[ImportFailure]) -> list[dict[str, Any]] | None:
        """失败明细入库，超出上限时截断（计数仍按真实值）。"""
        if not failures:
            return None
        limit = settings.QC_IMPORT_MAX_FAILURE_DETAIL
        return [
            {"row": item.row, "reason": item.reason, "raw": item.raw}
            for item in failures[:limit]
        ]

    def _mark_batch_failed(self, batch_id: int) -> None:
        self.db.rollback()
        batch = self.repo.get_batch(batch_id)
        if batch is None:  # pragma: no cover
            return
        batch.status = QcImportStatus.FAILED
        self.db.commit()
        logger.exception("对话导入写入失败：batch=%s", batch_id)


def _import_transcript_status(item: ParsedConversation) -> TranscriptStatus:
    """有录音的语音对话需转写，其余（纯文本）无需转写。"""
    if item.channel is ConversationChannel.VOICE and item.recording_path:
        return TranscriptStatus.PENDING
    return TranscriptStatus.NOT_REQUIRED


__all__ = [
    "ImportFailure",
    "ParsedConversation",
    "ParsedMessage",
    "QcConversationImportService",
    "SUPPORTED_EXTENSIONS",
    "detect_format",
    "parse_file",
]